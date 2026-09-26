"""把六道题的 Word 合并成一份提交用的报告：CM5235_HW2_Zhang_Xubo.docx（在 03_report_word/ 下）。

结构照 HW1：最前面一次大标题 + 姓名学号表，然后 Ex1…Ex6，每题之间分页。
合并用 docxcompose（不经过 Word COM——用 Word 合并多个 docx 会莫名卡死）。
用法：D:\\venvs\\cm5235_hw2\\Scripts\\python make_final_report.py
之后转 PDF：powershell -ExecutionPolicy Bypass -File docx_to_pdf.ps1 "<docx 完整路径>"
"""
import datetime
import sys

import docx
from docx.enum.text import WD_BREAK
from docxcompose.composer import Composer

from report_common import REPORT_DIR, add_runs, new_document, table

ORDER = [
    "HW2_Ex1_ground_state_2-chlorophenol.docx",
    "HW2_Ex2_theory_solvent_effects.docx",
    "HW2_Ex3_theory_resolution_of_identity.docx",
    "HW2_Ex4_excited_states_UV.docx",
    "HW2_Ex5_magnetic_exchange_BSDFT.docx",
    "HW2_Ex6_molecular_dynamics_water.docx",
]
STUDENT = [("Full Name", "Zhang Xubo"),
           ("Student Number", "A0359670R"),
           ("E-mail", "zhangxubo@u.nus.edu"),
           ("Date", datetime.date.today().strftime("%-d %B %Y") if sys.platform != "win32"
            else datetime.date.today().strftime("%d %B %Y").lstrip("0")),
           ("Signature", "")]


def build_cover():
    doc = new_document()
    p = doc.add_paragraph(style="Title")
    add_runs(p, "CM5235 - Homework ORCA (2026)")
    table(doc, [[k, v] for k, v in STUDENT], [4.0, 8.0], header_rows=0, bold_first_col=True,
          align_numbers=False)
    return doc


def main():
    missing = [f for f in ORDER if not (REPORT_DIR / f).exists()]
    if missing:
        sys.exit("缺这些题的 Word：" + str(missing))

    master = build_cover()
    composer = Composer(master)
    for f in ORDER:
        master.add_paragraph().add_run().add_break(WD_BREAK.PAGE)   # 每题另起一页
        composer.append(docx.Document(str(REPORT_DIR / f)))

    out = REPORT_DIR / "CM5235_HW2_Zhang_Xubo.docx"
    composer.save(str(out))
    print("写出", out)
    print("题目顺序：", ", ".join(f.split('_')[1] for f in ORDER))


if __name__ == "__main__":
    main()
