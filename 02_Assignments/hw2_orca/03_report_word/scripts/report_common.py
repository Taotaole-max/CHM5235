"""HW2 报告的共用排版工具：每道题的 make_exN_word.py 都 import 这里。

样式来源：HW1 的最终报告（Letter 纸、Calibri 10.5、Title/Heading 1/Heading 2、网格表 8.5 号字、
图注 9 号斜体居中）。第一次运行时从 HW1 的 docx 里清空正文，存成 hw1_style_template.docx，
之后每道题都从这个模板开始，格式和 HW1 一致。

文字里的上下标用简单标记：  cm^{-1}  →  上标；  E_{HOMO}  →  下标。
"""
import re
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

SCRIPTS = Path(__file__).resolve().parent
REPORT_DIR = SCRIPTS.parent                      # 03_report_word/
AS2 = REPORT_DIR.parent                          # as2/
TEMPLATE = SCRIPTS / "hw1_style_template.docx"
HW1_REPORT = AS2.parent / "as1" / "CM5235_HW1_Zhang_Xubo.docx"

# 动画放不进 PDF，和报告一起用邮件发给老师；文件名在报告里引用，复制到 as2/04_email_attachments/
GIF_EX1_ORCA12 = "HW2_ZhangXubo_Ex1_mode12_ORCA-numbering.gif"
GIF_EX1_AVO12 = "HW2_ZhangXubo_Ex1_mode12_Avogadro-numbering.gif"
GIF_EX6 = "HW2_ZhangXubo_Ex6_water_MD_XTB_Avogadro.gif"

TABLE_PT = 8.5
CAPTION_PT = 9


def _make_template():
    d = docx.Document(str(HW1_REPORT))
    body = d.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    d.save(str(TEMPLATE))


def new_document():
    if not TEMPLATE.exists():
        _make_template()
    return docx.Document(str(TEMPLATE))


# ------------------------------------------------------------------ 文字
_MARK = re.compile(r"(\^\{[^}]*\}|_\{[^}]*\})")


_NEG = re.compile(r"(?<![\w)])-(?=\d)")


def add_runs(paragraph, text, bold=None, italic=None, size=None):
    """按 ^{...} / _{...} 标记拆成普通、上标、下标几段。
    数字前的 ASCII 减号换成数学负号 −（U+2212）：Word 不会在它后面断行，也更规范。"""
    text = _NEG.sub("−", text)
    for piece in _MARK.split(text):
        if not piece:
            continue
        if piece.startswith("^{"):
            r = paragraph.add_run(piece[2:-1]); r.font.superscript = True
        elif piece.startswith("_{"):
            r = paragraph.add_run(piece[2:-1]); r.font.subscript = True
        else:
            r = paragraph.add_run(piece)
        if bold is not None:
            r.bold = bold
        if italic is not None:
            r.italic = italic
        if size is not None:
            r.font.size = Pt(size)
    return paragraph


def exercise_title(doc, text):
    """HW1 的写法：每道题开头一行加粗的 Normal 段落。"""
    return add_runs(doc.add_paragraph(), text, bold=True)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    return add_runs(p, text)


def para(doc, text, bold=False):
    return add_runs(doc.add_paragraph(), text, bold=bold or None)


def equation(doc, text):
    """单独成行的公式：居中。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return add_runs(p, text)


def para_lead(doc, lead, text):
    """HW1 理论题的写法：段首几个字加粗，后面接正文，同一段。"""
    p = doc.add_paragraph()
    add_runs(p, lead + " ", bold=True)
    return add_runs(p, text)


# ------------------------------------------------------------------ 表格
def table(doc, rows, col_widths_cm, header_rows=1, bold_first_col=True, merges=(), align_numbers=True):
    """rows：二维字符串列表（可含 ^{} _{} 标记）。merges：[(行, 起始列, 结束列), ...] 横向合并。"""
    ncol = len(col_widths_cm)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for (r, c0, c1) in merges:
        t.cell(r, c0).merge(t.cell(r, c1))
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            if text is None:
                continue
            cell = t.cell(i, j)
            cell.width = Cm(col_widths_cm[j])
            p = cell.paragraphs[0]
            bold = i < header_rows or (bold_first_col and j == 0)
            add_runs(p, str(text), bold=bold or None, size=TABLE_PT)
            if i < header_rows:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif align_numbers and j > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for j, w in enumerate(col_widths_cm):     # 列宽要在每个单元格和列上都设，Word 才认
        t.columns[j].width = Cm(w)
    for r in t.rows:                            # 表格不跨页拆行
        trPr = r._tr.get_or_add_trPr()
        cant = trPr.makeelement(qn("w:cantSplit"), {})
        trPr.append(cant)
    gap = doc.add_paragraph()                  # 表格和后面的正文之间留 4 磅
    gap.paragraph_format.space_after = Pt(0)
    gap.paragraph_format.line_spacing = Pt(4)
    return t


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(p, text, italic=True, size=CAPTION_PT)
    return p


def figure(doc, image_path, width_cm, caption_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True   # 图和图注不分页
    p.add_run().add_picture(str(image_path), width=Cm(width_cm))
    return caption(doc, caption_text)


def table_caption(doc, text):
    """表题放在表格上方，和下面的表格不分页。"""
    p = caption(doc, text)
    p.paragraph_format.keep_with_next = True
    return p


def save(doc, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    print("写出", path)
    return path
